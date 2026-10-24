import shutil

import pytest
from docx import Document
from docx.shared import Pt

import llm_client
import resume_customizer as rc
from conftest import make_resume


def fake_generate(system, prompt, chain, want_json=False, http=None):
    assert want_json and chain == llm_client.CUSTOMIZE_CHAIN
    assert "Spring Boot" in prompt  # the JD is sent
    return {
        "rewritten_summary": [{"id": 0, "text": "Java developer with 8 years building Spring Boot and Kafka "
                                                "microservices for banking clients."}],
        "rewritten_paragraphs": [
            {"id": 0, "text": "Built Spring Boot payment microservices on Kafka handling two million daily transactions."},
            {"id": 7, "text": "unknown id is ignored"},
            {"id": 1, "text": ""},
        ],
        "rewritten_skills": [{"id": 0, "text": "Languages: Java, SQL, Spring Boot, Kafka"}],
    }


JD = "Senior Java Developer, Spring Boot, Kafka, AWS, C2C Dallas TX"


def test_extract_structure_finds_sections(tmp_path):
    doc = Document(str(make_resume(tmp_path / "r.docx")))
    s = rc.extract_structure(doc)
    assert [p.text for p in s.summary] == [
        "Java developer with 8 years building enterprise applications for banking clients."]
    # the short "Client: ... | dates" line is kept out of the rewrite set
    assert [p.text[:20] for p in s.paragraphs] == ["Built payment micros", "Led migration of leg"]
    assert [p.text for p in s.skills] == ["Languages: Java, SQL"]


def test_tailor_rewrites_and_preserves_formatting(tmp_path):
    doc = Document(str(make_resume(tmp_path / "r.docx")))
    changed = rc.tailor(doc, JD, "Java Developer", generate=fake_generate)
    assert changed == 3
    s = rc.extract_structure(doc)
    run = s.summary[0].runs[0]
    assert "Spring Boot and Kafka" in s.summary[0].text
    assert (run.font.name, run.font.size, run.italic) == ("Calibri", Pt(11), True)
    assert str(run.font.color.rgb) == "223344"
    assert s.paragraphs[0].text.startswith("Built Spring Boot payment")
    assert s.paragraphs[1].text.startswith("Led migration")  # untouched (empty rewrite rejected)
    skills = s.skills[0]
    assert skills.text == "Languages: Java, SQL, Spring Boot, Kafka"
    assert skills.runs[0].bold and skills.runs[0].text == "Languages:" and not skills.runs[1].bold
    assert "Client: Acme Bank" in "\n".join(p.text for p in doc.paragraphs)


def test_remove_consecutive_blanks(tmp_path):
    doc = Document(str(make_resume(tmp_path / "r.docx")))
    assert rc.remove_consecutive_blanks(doc) == 1


def test_customize_falls_back_to_original_when_ai_unavailable(workspace):
    def failing(*a, **k):
        raise llm_client.LLMError("No API key set")

    resume = workspace / "jane_resume.docx"
    result = rc.customize(resume, JD, "Java Developer", {"name": "Jane Doe"}, generate=failing)
    assert not result.customized and "AI unavailable" in result.note
    assert result.path.exists()


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_customize_produces_real_pdf(workspace):
    from pypdf import PdfReader

    result = rc.customize(workspace / "jane_resume.docx", JD, "Java Developer", {"name": "Jane Doe"},
                          generate=fake_generate)
    assert result.customized and result.path.suffix == ".pdf"
    text = "".join(page.extract_text() for page in PdfReader(str(result.path)).pages)
    assert "Spring Boot payment microservices" in text and "Acme Bank" in text


def test_pdf_resume_roundtrip(workspace):
    pytest.importorskip("pdf2docx")
    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("needs LibreOffice to build the PDF input")
    pdf = rc.docx_to_pdf(workspace / "jane_resume.docx")
    result = rc.customize(pdf, JD, "Java Developer", {"name": "Jane Doe"}, generate=fake_generate)
    assert result.path.exists()


def test_parse_json_tolerates_fences():
    assert llm_client.parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert llm_client.parse_json('Here you go: {"a": 2} thanks') == {"a": 2}
