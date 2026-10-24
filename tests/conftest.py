import sys
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import client_config  # noqa: E402
import llm_client  # noqa: E402


@pytest.fixture(autouse=True)
def fresh_ai_state(monkeypatch):
    """No discovered models, failure counts or cool-downs carry over between tests; no real sleeps."""
    llm_client.reset()
    monkeypatch.setattr(llm_client, "_sleep", lambda s: None)
    yield
    llm_client.reset()

FIXTURES = Path(__file__).resolve().parent / "fixtures"

ENV = {
    "LINKEDIN_EMAIL": "operator@example.com",
    "LINKEDIN_PASSWORD": "correct-password",
    "GMAIL_EMAIL": "operator@gmail.com",
    "GMAIL_APP_PASSWORD": "abcd efgh ijkl mnop",
    "GEMINI_API_KEY": "", "GROQ_API_KEY": "", "OPENROUTER_API_KEY": "",
}


def make_resume(path: Path) -> Path:
    doc = Document()
    head = doc.add_paragraph()
    head.add_run("PROFESSIONAL SUMMARY").bold = True
    summary = doc.add_paragraph()
    run = summary.add_run("Java developer with 8 years building enterprise applications for banking clients.")
    run.font.name, run.font.size, run.italic = "Calibri", Pt(11), True
    run.font.color.rgb = RGBColor(0x22, 0x33, 0x44)
    doc.add_paragraph("")
    doc.add_paragraph("")
    doc.add_paragraph().add_run("PROFESSIONAL EXPERIENCE").bold = True
    doc.add_paragraph("Client: Acme Bank, Dallas TX  |  Jan 2022 - Present")
    doc.add_paragraph("Built payment microservices handling two million daily transactions with high availability.",
                      style="List Bullet")
    doc.add_paragraph("Led migration of legacy monolith modules to a cloud-hosted service architecture.",
                      style="List Bullet")
    doc.add_paragraph().add_run("TECHNICAL SKILLS").bold = True
    table = doc.add_table(rows=1, cols=1)
    cell_p = table.rows[0].cells[0].paragraphs[0]
    cell_p.add_run("Languages:").bold = True
    cell_p.add_run(" Java, SQL")
    doc.add_paragraph().add_run("EDUCATION").bold = True
    doc.add_paragraph("B.Tech in Computer Science, 2015")
    doc.save(str(path))
    return path


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    """Isolated DATA_DIR, env vars and a single test candidate with a real DOCX resume."""
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    for k, v in ENV.items():
        monkeypatch.setenv(k, v)
    resume = make_resume(tmp_path / "jane_resume.docx")
    monkeypatch.setattr(client_config, "CANDIDATE_ROLES", {
        "jane": {"primary_role": "Java Developer", "search_roles": ["Java Developer"],
                 "allowed_domains": ["java", "fullstack", "backend"],
                 "exclusion_keywords": client_config.DEFAULT_EXCLUSION_KEYWORDS},
    })
    monkeypatch.setattr(client_config, "CLIENTS", {
        "jane": {"name": "Jane Doe", "email": "jane@candidate.com", "phone": "555-019-2834",
                 "linkedin": "https://www.linkedin.com/in/janedoe", "location": "Dallas, TX",
                 "relocation": "Yes", "work_auth": "H1B", "availability": "Immediate",
                 "experience": "8+ Years", "resume_candidates": [str(resume)], "customizer": "default"},
    })
    monkeypatch.setattr(client_config, "TEAM_LEAD_EMAIL", "lead@agency.com")
    monkeypatch.setattr(client_config, "ADMIN_EMAIL", "admin@agency.com")
    return tmp_path
