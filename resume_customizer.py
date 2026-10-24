"""AI resume customizer: parse a DOCX/PDF resume, tailor summary/projects/skills to a JD, keep styling, export PDF."""

from __future__ import annotations

import contextlib
import io
import json
import logging
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

import llm_client
import settings
from lock_util import word_com_lock

log = logging.getLogger(__name__)

SUMMARY_HEADERS = {"professional summary", "summary", "profile summary", "career summary",
                   "professional profile", "profile", "career objective", "objective", "about me"}
PROJECT_HEADERS = {"professional experience", "work experience", "experience", "project experience",
                   "projects", "employment history", "relevant experience", "project details"}
SKILLS_HEADERS = {"technical skills", "skills", "core competencies", "technical expertise", "skill set",
                  "technical summary", "key skills"}
OTHER_HEADERS = {"education", "certifications", "certification", "awards", "achievements", "languages",
                 "references", "personal details", "declaration", "publications", "training", "education details"}
ALL_HEADERS = SUMMARY_HEADERS | PROJECT_HEADERS | SKILLS_HEADERS | OTHER_HEADERS

SYSTEM_PROMPT = (
    "You tailor IT consultant resumes to a specific job description for US C2C contract roles. "
    "Rewrite the given summary paragraphs, project/experience paragraphs and skills lines so the key "
    "technologies and responsibilities from the job description appear naturally, weaving the most important "
    "JD technologies 2 to 3 times per project where they fit the original work. Keep each item's historical "
    "context: never change employers, clients, job titles, dates, degrees or certifications, and keep each "
    "rewritten item about the same length as the original. If a skills line starts with a label such as "
    "'Languages:', keep that label. Return JSON only."
)


@dataclass
class ResumeStructure:
    summary: list[Paragraph] = field(default_factory=list)
    paragraphs: list[Paragraph] = field(default_factory=list)
    skills: list[Paragraph] = field(default_factory=list)


@dataclass
class CustomizationResult:
    path: Path
    customized: bool
    note: str = ""


# ------------------------------------------------------------------ parsing

def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[:\-–|•*]+$", "", text.strip())).lower().replace(" & ", " and ")


def header_kind(paragraph: Paragraph) -> str | None:
    """'summary' | 'projects' | 'skills' | 'other' if this paragraph is a section heading."""
    text = _norm(paragraph.text)
    if not text or len(text) > 40:
        return None
    if text in SUMMARY_HEADERS:
        return "summary"
    if text in PROJECT_HEADERS:
        return "projects"
    if text in SKILLS_HEADERS:
        return "skills"
    if text in OTHER_HEADERS:
        return "other"
    # Only known section names count: company/client lines are often styled as headings too.
    return None


def iter_block_paragraphs(doc):
    """Body paragraphs and table-cell paragraphs in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, doc)
        elif tag == "tbl":
            table = Table(child, doc)
            seen = set()
            for row in table.rows:
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    yield from cell.paragraphs


def _is_list_item(p: Paragraph) -> bool:
    style = (p.style.name or "").lower() if p.style is not None else ""
    has_numbering = p._p.pPr is not None and p._p.pPr.numPr is not None
    return has_numbering or "list" in style or p.text.lstrip()[:1] in "•●▪-–*"


def extract_structure(doc) -> ResumeStructure:
    structure = ResumeStructure()
    section = None
    for p in iter_block_paragraphs(doc):
        kind = header_kind(p)
        if kind:
            section = kind
            continue
        text = p.text.strip()
        if not text:
            continue
        if section == "summary":
            structure.summary.append(p)
        elif section == "projects" and (len(text) >= 60 or _is_list_item(p)):
            # Short lines (company, client, dates, title) are kept untouched.
            structure.paragraphs.append(p)
        elif section == "skills":
            structure.skills.append(p)
    return structure


# ------------------------------------------------------------------ writing

def set_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    """Replace text while keeping the original run formatting (font, size, bold, colour, style)."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return
    first = runs[0].text
    label = first.strip()
    if len(runs) > 1 and label.endswith(":") and new_text.strip().startswith(label):
        runs[0].text = first
        runs[1].text = new_text.strip()[len(label):]
        for run in runs[2:]:
            run.text = ""
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def remove_consecutive_blanks(doc) -> int:
    removed = 0
    previous_blank = False
    for p in list(doc.paragraphs):
        blank = not p.text.strip() and not p._p.xpath(".//w:drawing | .//w:pict")
        if blank and previous_blank:
            p._p.getparent().remove(p._p)
            removed += 1
            continue
        previous_blank = blank
    return removed


def _items(paragraphs: list[Paragraph]) -> list[dict]:
    return [{"id": i, "text": p.text.strip()} for i, p in enumerate(paragraphs)]


def apply_rewrites(paragraphs: list[Paragraph], rewrites, label: str) -> int:
    """Write validated rewrites back. Ignores unknown ids, empty text and wildly longer output."""
    applied = 0
    for item in rewrites or []:
        try:
            idx, text = int(item["id"]), str(item["text"]).strip()
        except (KeyError, TypeError, ValueError):
            continue
        if not (0 <= idx < len(paragraphs)) or not text:
            continue
        original = paragraphs[idx].text.strip()
        if len(text) > 3 * len(original) + 300:
            log.warning("Skipping %s #%d: rewrite is far longer than the original", label, idx)
            continue
        set_paragraph_text(paragraphs[idx], text)
        applied += 1
    return applied


def tailor(doc, job_description: str, role: str, generate=llm_client.generate, parser=None) -> int:
    """Send the resume structure to the LLM chain and apply the rewrites. Returns paragraphs changed."""
    structure = (parser or extract_structure)(doc)
    if not (structure.summary or structure.paragraphs or structure.skills):
        raise llm_client.LLMError("Could not find summary/experience/skills sections in the resume")
    payload = {
        "target_role": role,
        "job_description": job_description[:6000],
        "summary": _items(structure.summary),
        "paragraphs": _items(structure.paragraphs),
        "skills": _items(structure.skills),
    }
    prompt = (
        "Tailor this resume content to the job description.\n\n"
        f"{json.dumps(payload, ensure_ascii=False)}\n\n"
        'Respond with JSON: {"rewritten_summary": [{"id": int, "text": str}], '
        '"rewritten_paragraphs": [{"id": int, "text": str}], "rewritten_skills": [{"id": int, "text": str}]} '
        "using the same ids as the input items."
    )
    result = generate(SYSTEM_PROMPT, prompt, llm_client.CUSTOMIZE_CHAIN, want_json=True)
    return (apply_rewrites(structure.summary, result.get("rewritten_summary"), "summary")
            + apply_rewrites(structure.paragraphs, result.get("rewritten_paragraphs"), "paragraph")
            + apply_rewrites(structure.skills, result.get("rewritten_skills"), "skills"))


# ------------------------------------------------------------------ conversion

def pdf_to_docx(pdf_path: Path, out_dir: Path) -> Path:
    from pdf2docx import Converter

    docx_path = out_dir / f"{pdf_path.stem}.docx"
    converter = Converter(str(pdf_path))
    try:
        converter.convert(str(docx_path))
    finally:
        converter.close()
    return docx_path


def _kill_word() -> None:
    if sys.platform == "win32":
        subprocess.run(["taskkill", "/f", "/im", "winword.exe"], capture_output=True)


def docx_to_pdf(docx_path: Path) -> Path | None:
    """MS Word COM via docx2pdf (Windows/macOS) under a mutex; LibreOffice elsewhere. None if impossible."""
    pdf_path = docx_path.with_suffix(".pdf")
    if sys.platform in ("win32", "darwin"):
        for attempt in (1, 2):
            try:
                from docx2pdf import convert

                with word_com_lock(), contextlib.redirect_stderr(io.StringIO()):  # hide its progress bar
                    convert(str(docx_path), str(pdf_path))
                if pdf_path.exists():
                    return pdf_path
            except Exception as exc:
                log.warning("Word PDF conversion failed (attempt %d): %s", attempt, exc)
                _kill_word()
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        try:
            with word_com_lock():
                subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent),
                                str(docx_path)], capture_output=True, timeout=180)
        except (subprocess.SubprocessError, OSError) as exc:
            log.warning("LibreOffice PDF conversion failed: %s", exc)
        if pdf_path.exists():
            return pdf_path
    log.warning("No PDF converter available (install MS Word or LibreOffice); sending DOCX")
    return None


def ensure_pdf(resume_path: Path) -> Path:
    """Return a PDF version of the resume (converting a DOCX copy if needed), or the original file."""
    if resume_path.suffix.lower() == ".pdf":
        return resume_path
    work = settings.scratch_dir() / "originals"
    work.mkdir(parents=True, exist_ok=True)
    copy = work / resume_path.name
    shutil.copy2(resume_path, copy)
    return docx_to_pdf(copy) or copy


def _slug(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_") or "resume"


def customize(resume_path: Path, job_description: str, role: str, candidate: dict,
              generate=llm_client.generate, parser=None) -> CustomizationResult:
    """Tailor the candidate's resume to one JD. Falls back to the original resume if the AI step fails."""
    work = settings.scratch_dir() / f"{_slug(candidate.get('name', 'candidate'))}_{datetime.now():%Y%m%d_%H%M%S_%f}"
    work.mkdir(parents=True, exist_ok=True)
    source = pdf_to_docx(resume_path, work) if resume_path.suffix.lower() == ".pdf" else resume_path

    doc = Document(str(source))
    note = ""
    try:
        changed = tailor(doc, job_description, role, generate, parser)
        note = f"{changed} sections rewritten"
    except llm_client.LLMError as exc:
        changed = 0
        note = "original resume (AI unavailable)"
        log.warning("AI tailoring skipped: %s", exc)
    remove_consecutive_blanks(doc)

    out_docx = work / f"{_slug(candidate.get('name', 'candidate'))}_Resume.docx"
    doc.save(str(out_docx))
    pdf = docx_to_pdf(out_docx)
    return CustomizationResult(path=pdf or out_docx, customized=changed > 0, note=note)
